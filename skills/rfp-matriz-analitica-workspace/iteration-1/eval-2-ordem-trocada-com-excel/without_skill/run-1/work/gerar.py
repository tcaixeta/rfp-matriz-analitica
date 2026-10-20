# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dados import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)

# ============================ EXCEL ============================
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
SEV = {"Crítica": "C00000", "Crítico": "C00000", "Alta": "F4B183", "Alto": "F4B183",
       "Média": "FFE699", "Médio": "FFE699", "Baixa": "C6E0B4", "Baixo": "C6E0B4"}

wb = Workbook()

def sheet(title, headers, rows, widths, sev_cols=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in rows:
        ws.append(list(row))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    last = f"{get_column_letter(len(headers))}{ws.max_row}"
    ws.auto_filter.ref = f"A1:{last}"
    ws.row_dimensions[1].height = 32
    for col in sev_cols:
        L = get_column_letter(col)
        rng = f"{L}2:{L}{ws.max_row}"
        for k, color in SEV.items():
            font = Font(bold=True, color="FFFFFF") if color == "C00000" else Font(bold=True)
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{L}2="{k}"'], fill=PatternFill("solid", fgColor=color), font=font))
    ws.sheet_view.zoomScale = 90
    return ws

# --- Resumo
ws = wb.active
ws.title = "Resumo"
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 150
lines = [("Matriz de Análise - RFP nº 2026-017 - Banco Horizonte Fictício S.A. - Serviços gerenciados de SOC 24x7 e IAM", "t"),
         ("Documentos fictícios para teste. Análise baseada no RFP v1.0, Anexo I v1.0, apresentação do workshop e Q&A Rodada 1.", "i"),
         ("", ""), ("RECOMENDAÇÃO", "h"), (RECOMENDACAO, "p"), ("", ""), ("PONTOS PRINCIPAIS", "h")]
lines += [(f"{i}. {t}", "p") for i, t in enumerate(TOP_PONTOS, 1)]
lines += [("", ""), ("HIERARQUIA DOCUMENTAL ADOTADA", "h")] + [(t, "p") for t in HIERARQUIA]
lines += [("", ""), ("COMO USAR ESTA PLANILHA", "h"),
          ("Matriz: todos os requisitos e cláusulas. Filtre as colunas 'Jurídico?' ou 'Pricing?' = S para separar o que vai para cada área. Use 'Criticidade' para priorizar. As colunas 'Responsável', 'Status de tratamento' e 'Notas' estão livres para o seu acompanhamento.", "p"),
          ("Conflitos: divergências entre documentos e a leitura adotada. Jurídico: cláusulas de risco e contrapropostas. Pricing: drivers e premissas de custo. Perguntas: esclarecimentos para enviar ao Banco ou registrar como premissas. Questionário: resposta sugerida para QT-01 a QT-06 com os pesos da aba oculta. Cronograma e Documentos: datas vigentes, fontes e lacunas.", "p"),
          ("Referências: RFP = PDF principal; AI = Anexo I (docx); PPT = apresentação do workshop; XLSX = Q&A/Questionário; Q&A nº X = resposta da Rodada 1.", "p")]
for i, (t, k) in enumerate(lines, 1):
    c = ws.cell(row=i, column=2, value=t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if k == "t": c.font = Font(bold=True, size=14, color="1F3864")
    elif k == "h": c.font = Font(bold=True, size=11, color="FFFFFF"); c.fill = HDR_FILL
    elif k == "i": c.font = Font(italic=True, color="595959")
ws.sheet_view.showGridLines = False

# --- Matriz
mrows = []
for (i, fonte, item, tema, req, tipo, sit, alt, crit, jur, pri, area, imp, rec) in M:
    mrows.append((i, fonte, item, tema, req, tipo, sit, alt, crit, jur, pri, area, imp, rec, "", "Pendente", ""))
ws = sheet("Matriz", ["ID", "Documento", "Item", "Tema", "Requisito / cláusula (resumo)", "Tipo", "Situação vigente (após Q&A e conflitos)",
                      "Alterado por", "Criticidade", "Jurídico?", "Pricing?", "Área líder", "Impacto em preço", "Recomendação / ação",
                      "Responsável", "Status de tratamento", "Notas"],
           mrows, [6, 10, 14, 15, 55, 16, 34, 11, 11, 9, 9, 16, 11, 60, 14, 14, 25], sev_cols=(9, 13))
dv = DataValidation(type="list", formula1='"Pendente,Em análise,Enviado ao Banco,Resolvido,Aceito como risco"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"P2:P{ws.max_row}")
dvsn = DataValidation(type="list", formula1='"S,N"', allow_blank=True)
ws.add_data_validation(dvsn); dvsn.add(f"J2:K{ws.max_row}")

sheet("Conflitos", ["ID", "Tema", "Fonte A", "Fonte B", "Descrição", "Resolvido pelo Q&A?", "Leitura adotada", "Ação"],
      CONFLITOS, [6, 20, 32, 36, 50, 13, 36, 50])
sheet("Jurídico", ["ID", "Cláusula", "Texto (resumo)", "Risco", "Severidade", "Posição recomendada / contraproposta"],
      JURIDICO, [6, 30, 45, 45, 11, 70], sev_cols=(5,))
sheet("Pricing", ["ID", "Driver de custo", "Valor / premissa", "Fonte", "Impacto", "Observação"],
      PRICING, [6, 30, 40, 30, 10, 70], sev_cols=(5,))
sheet("Perguntas", ["ID", "Referência", "Pergunta de esclarecimento", "Motivo", "Área"], PERGUNTAS, [6, 20, 80, 35, 18])
sheet("Questionário", ["ID", "Requisito", "Obrigatório", "Peso (aba oculta)", "Referência cruzada", "Resposta sugerida", "Comentário / evidência"],
      QUESTIONARIO, [7, 50, 11, 10, 22, 30, 60])
sheet("Cronograma", ["Marco", "Data original", "Data vigente", "Fonte", "Observação"], CRONOGRAMA, [36, 20, 32, 22, 70])
docs_rows = [(s, a, d, v, o) for (s, a, d, v, o) in DOCS] + [("LACUNA", n, d, "-", a) for (n, d, a) in LACUNAS]
sheet("Documentos", ["Sigla", "Arquivo / item", "Descrição", "Versão / data", "Observações / ação"], docs_rows, [9, 38, 45, 20, 80])

xlsx_path = os.path.join(OUT, "Matriz_Analise_RFP_2026-017.xlsx")
wb.save(xlsx_path)

# ============================ WORD ============================
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.8))

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
    tcPr.append(s)

def table(headers, rows, widths=None, sev_idx=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        run = c.paragraphs[0].add_run(h); run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255); run.font.size = Pt(9)
        shade(c, "1F3864")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v)); run.font.size = Pt(8.5)
            if sev_idx is not None and i == sev_idx and str(v) in SEV:
                shade(cells[i], SEV[str(v)])
                run.bold = True
                if SEV[str(v)] == "C00000": run.font.color.rgb = RGBColor(255, 255, 255)
    # repeat header
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

title = doc.add_heading("Matriz de Análise - RFP nº 2026-017", 0)
p = doc.add_paragraph()
p.add_run("Banco Horizonte Fictício S.A. - Serviços gerenciados de SOC 24x7x365 e Gestão de Identidades e Acessos (IAM)\n").bold = True
p.add_run("Base: RFP v1.0 (15/10/2026), Anexo I v1.0, apresentação do workshop (22/10/2026) e Q&A Rodada 1 (12/11/2026). "
          "Documentos fictícios para teste. A versão filtrável está em Matriz_Analise_RFP_2026-017.xlsx.").italic = True

doc.add_heading("1. Recomendação", 1)
doc.add_paragraph(RECOMENDACAO)
doc.add_heading("2. Pontos principais", 1)
for t in TOP_PONTOS:
    doc.add_paragraph(t, style="List Number")

doc.add_heading("3. Documentos analisados e hierarquia", 1)
doc.add_paragraph("Os arquivos chegaram fora de ordem (Q&A primeiro, RFP por último). A análise foi reorganizada pela "
                  "hierarquia abaixo, independentemente da ordem de envio.")
table(["Sigla", "Arquivo", "Descrição", "Versão / data", "Observações de leitura"], DOCS, [1.5, 6, 5, 3, 10])
for t in HIERARQUIA:
    doc.add_paragraph(t, style="List Bullet")

doc.add_heading("4. Cronograma vigente", 1)
table(["Marco", "Data original", "Data vigente", "Fonte", "Observação"], CRONOGRAMA, [5, 3, 5, 3.5, 9])

doc.add_heading("5. Conflitos e inconsistências entre documentos", 1)
table(["ID", "Tema", "Fonte A", "Fonte B", "Descrição", "Resolvido?", "Leitura adotada", "Ação"], CONFLITOS,
      [1, 2.5, 3.5, 3.5, 5, 1.8, 3.5, 4.5])

doc.add_heading("6. Matriz de requisitos", 1)
doc.add_paragraph("J = encaminhar ao Jurídico; P = encaminhar ao Pricing. Colunas de acompanhamento estão no Excel.")
table(["ID", "Doc / item", "Requisito / cláusula", "Tipo", "Situação vigente", "Crit.", "J", "P", "Recomendação"],
      [(m[0], f"{m[1]} {m[2]}", m[4], m[5], m[6] + ("" if m[7] == "-" else f" ({m[7]})"), m[8], m[9], m[10], m[13]) for m in M],
      [1, 2.4, 6, 2.2, 3.8, 1.5, 0.7, 0.7, 7], sev_idx=5)

doc.add_heading("7. Riscos jurídicos e contrapropostas", 1)
table(["ID", "Cláusula", "Texto", "Risco", "Severidade", "Posição recomendada"], JURIDICO, [1, 3.5, 5, 5, 1.8, 9], sev_idx=4)

doc.add_heading("8. Drivers e premissas de pricing", 1)
table(["ID", "Driver", "Valor / premissa", "Fonte", "Impacto", "Observação"], PRICING, [1, 4, 5, 4, 1.6, 9], sev_idx=4)

doc.add_heading("9. Questionário técnico (com pesos da aba oculta)", 1)
doc.add_paragraph("A aba 'Pesos', oculta no XLSX do Banco, define pesos por item e penaliza com -2 pontos cada requisito "
                  "não obrigatório respondido 'Não atende'. Os itens não obrigatórios são QT-02 e QT-05.")
table(["ID", "Requisito", "Obrig.", "Peso", "Ref.", "Resposta sugerida", "Evidência"], QUESTIONARIO, [1.3, 6, 1.3, 1.1, 3, 4.5, 8])

doc.add_heading("10. Perguntas de esclarecimento", 1)
doc.add_paragraph("O prazo de perguntas do cronograma original era 05/11/2026. Se não houver nova rodada, cada pergunta "
                  "abaixo deve virar premissa ou ressalva expressa na proposta.")
table(["ID", "Referência", "Pergunta", "Motivo", "Área"], PERGUNTAS, [1, 3, 11, 5, 3])

doc.add_heading("11. Lacunas documentais", 1)
table(["Item", "Situação", "Ação"], LACUNAS, [6, 10, 9])

docx_path = os.path.join(OUT, "Matriz_Analise_RFP_2026-017.docx")
doc.save(docx_path)
print(xlsx_path); print(docx_path)
