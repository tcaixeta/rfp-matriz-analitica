import sys, pdfplumber, docx, pptx, openpyxl
d, w = sys.argv[1], sys.argv[2]
with pdfplumber.open(f"{d}/RFP-2026-017_Servicos_SOC_IAM.pdf") as pdf:
    out=[]
    print("PDF pages", len(pdf.pages), pdf.metadata)
    for i,p in enumerate(pdf.pages,1):
        out.append(f"\n===== PAGINA {i} =====\n"+(p.extract_text() or ""))
        for t in p.extract_tables():
            out.append("TABELA: "+repr(t))
    open(f"{w}/rfp.txt","w").write("\n".join(out))
doc = docx.Document(f"{d}/Anexo_I_Especificacao_Tecnica.docx")
out=[]
cp=doc.core_properties; out.append(f"PROPS: title={cp.title} author={cp.author} created={cp.created} modified={cp.modified} rev={cp.revision}")
from docx.oxml.ns import qn
body=doc.element.body
for el in body.iterchildren():
    if el.tag==qn('w:p'):
        p=docx.text.paragraph.Paragraph(el,doc); 
        if p.text.strip(): out.append(f"[{p.style.name}] {p.text}")
    elif el.tag==qn('w:tbl'):
        t=docx.table.Table(el,doc)
        out.append("TABELA:")
        for r in t.rows: out.append(" | ".join(c.text for c in r.cells))
for s in doc.sections:
    for p in s.header.paragraphs+s.footer.paragraphs:
        if p.text.strip(): out.append("HDR/FTR: "+p.text)
# comments / tracked changes
import zipfile,re
z=zipfile.ZipFile(f"{d}/Anexo_I_Especificacao_Tecnica.docx")
out.append("ZIP: "+", ".join(z.namelist()))
x=z.read("word/document.xml").decode()
for tag in ["w:ins","w:del","w:commentRangeStart","w:vanish"]:
    out.append(f"{tag}: {x.count('<'+tag)}")
open(f"{w}/anexo1.txt","w").write("\n".join(out))
pr=pptx.Presentation(f"{d}/Apresentacao_Ambiente_Atual.pptx")
out=[]
cp=pr.core_properties; out.append(f"PROPS: title={cp.title} author={cp.author} created={cp.created} modified={cp.modified}")
for i,s in enumerate(pr.slides,1):
    out.append(f"\n===== SLIDE {i} =====")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip(): out.append(sh.text_frame.text)
        if sh.has_table:
            for r in sh.table.rows: out.append(" | ".join(c.text for c in r.cells))
        if sh.shape_type==13: out.append("[IMAGEM]")
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        out.append("NOTAS: "+s.notes_slide.notes_text_frame.text)
open(f"{w}/pptx.txt","w").write("\n".join(out))
wb=openpyxl.load_workbook(f"{d}/QA_Rodada1_e_Questionario.xlsx")
out=[]
for ws in wb.worksheets:
    out.append(f"\n===== ABA {ws.title} ({ws.sheet_state}) dims={ws.dimensions} =====")
    for r in ws.iter_rows():
        vals=[(c.coordinate,c.value) for c in r if c.value is not None]
        if vals: out.append(" | ".join(f"{k}={v}" for k,v in vals))
    for c in [c for r in ws.iter_rows() for c in r if c.comment]: out.append(f"COMENTARIO {c.coordinate}: {c.comment.text}")
open(f"{w}/xlsx.txt","w").write("\n".join(out))
