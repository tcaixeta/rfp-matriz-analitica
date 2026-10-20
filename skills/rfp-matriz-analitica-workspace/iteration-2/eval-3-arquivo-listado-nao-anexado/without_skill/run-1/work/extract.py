import sys, pdfplumber, docx, pptx, openpyxl, os
D=sys.argv[1]; W=sys.argv[2]
with pdfplumber.open(os.path.join(D,"RFP-2026-017_Servicos_SOC_IAM.pdf")) as pdf:
    out=[]
    for i,p in enumerate(pdf.pages,1):
        out.append(f"=== PÁGINA {i} ===\n"+(p.extract_text() or ""))
        for t in p.extract_tables(): out.append("TABELA: "+repr(t))
    open(os.path.join(W,"rfp.txt"),"w").write("\n".join(out))
d=docx.Document(os.path.join(D,"Anexo_I_Especificacao_Tecnica.docx"))
out=[]
from docx.oxml.ns import qn
body=d.element.body
for el in body.iterchildren():
    if el.tag==qn('w:p'):
        p=docx.text.paragraph.Paragraph(el,d)
        if p.text.strip(): out.append(f"[{p.style.name}] {p.text}")
    elif el.tag==qn('w:tbl'):
        t=docx.table.Table(el,d)
        for r in t.rows: out.append("| "+" | ".join(c.text for c in r.cells)+" |")
for s in d.sections:
    for p in s.header.paragraphs:
        if p.text.strip(): out.append("[HEADER] "+p.text)
    for p in s.footer.paragraphs:
        if p.text.strip(): out.append("[FOOTER] "+p.text)
cp=d.core_properties; out.append(f"[CORE] title={cp.title} author={cp.author} created={cp.created} modified={cp.modified}")
open(os.path.join(W,"anexo1.txt"),"w").write("\n".join(out))
pr=pptx.Presentation(os.path.join(D,"Apresentacao_Ambiente_Atual.pptx"))
out=[]
for i,s in enumerate(pr.slides,1):
    out.append(f"=== SLIDE {i} ===")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip(): out.append(sh.text_frame.text)
        if sh.has_table:
            for r in sh.table.rows: out.append("| "+" | ".join(c.text for c in r.cells)+" |")
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(): out.append("[NOTAS] "+s.notes_slide.notes_text_frame.text)
open(os.path.join(W,"pptx.txt"),"w").write("\n".join(out))
wb=openpyxl.load_workbook(os.path.join(D,"QA_Rodada1_e_Questionario.xlsx"))
out=[]
for ws in wb.worksheets:
    out.append(f"=== ABA {ws.title} ({ws.sheet_state}) dims={ws.dimensions} ===")
    for row in ws.iter_rows():
        vals=[(c.coordinate,c.value) for c in row if c.value is not None]
        if vals: out.append(" ; ".join(f"{k}={v}" for k,v in vals))
    for c in [c for row in ws.iter_rows() for c in row if c.comment]: out.append(f"[COMENT {c.coordinate}] {c.comment.text}")
open(os.path.join(W,"xlsx.txt"),"w").write("\n".join(out))
