import sys, docx
from docx.document import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
d=docx.Document(sys.argv[1])
print('CORE:',d.core_properties.title,d.core_properties.author,d.core_properties.comments)
body=d.element.body
for ch in body.iterchildren():
    tag=ch.tag.split('}')[1]
    if tag=='p':
        p=Paragraph(ch,d)
        t=p.text
        if t.strip(): print(f'[{p.style.name}] {t}')
    elif tag=='tbl':
        tb=Table(ch,d)
        for r in tb.rows: print('  | '+' | '.join(c.text for c in r.cells))
for s in d.sections:
    for p in s.header.paragraphs:
        if p.text.strip(): print('HEADER:',p.text)
    for p in s.footer.paragraphs:
        if p.text.strip(): print('FOOTER:',p.text)
