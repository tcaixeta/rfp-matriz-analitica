# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dados import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
RISK_FILL = {"Alto": "F8CBAD", "Médio": "FFE699", "Baixo": "C6EFCE"}

def sheet_table(ws, headers, rows, widths, name, start_row=1, risk_col=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = WRAP; c.border = BORDER
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    last = start_row + len(rows)
    ref = f"A{start_row}:{get_column_letter(len(headers))}{last}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3 if len(headers) > 4 else 2)
    if risk_col:
        col = get_column_letter(risk_col)
        rng = f"{col}{start_row+1}:{col}{last}"
        for k, color in RISK_FILL.items():
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{start_row+1}="{k}"'], fill=PatternFill("solid", fgColor=color)))
    return last

wb = Workbook()

# ---------- Leia-me ----------
ws = wb.active; ws.title = "Leia-me"
ws["A1"] = "Matriz de Análise — RFP nº 2026-017 — Banco Horizonte Fictício S.A. — SOC 24x7 + IAM"
ws["A1"].font = Font(bold=True, size=14, color="1F3864")
ws["A2"] = "Data da análise: 18/10/2026 · Entrega das propostas: 27/11/2026 até 18h (prorrogada pelo Q&A #1) · Perguntas até 05/11/2026"
ws["A3"] = "DOCUMENTOS DE ORIGEM FICTÍCIOS PARA TESTE."
r = 5
ws.cell(row=r, column=1, value="Como usar").font = Font(bold=True); r += 1
for t in [
    "• 'Matriz': todos os itens analisados, com filtros. Para o JURÍDICO, filtrar 'Impacto Jurídico' = Sim. Para o PRICING, filtrar 'Impacto Pricing' = Sim. As abas 'Jurídico' e 'Pricing' já trazem esses recortes prontos.",
    "• 'Conflitos': divergências entre documentos e a posição adotada.",
    "• 'Esclarecimentos': perguntas prontas para enviar até 05/11/2026 (colunas 'Status' e 'Resposta do banco' para acompanhamento).",
    "• 'Pontuação' e 'Questionário': simulação da nota técnica e recomendação de resposta para QT-01 a QT-06.",
    "• 'Cronograma': marcos oficiais e datas derivadas.",
    "• A coluna 'Status interno' da Matriz está livre para uso das áreas (lista: Aberto / Em análise / Resolvido / Aceito como risco).",
]:
    ws.cell(row=r, column=1, value=t); r += 1
r += 1
ws.cell(row=r, column=1, value="Documentos recebidos (na ordem enviada) e hierarquia adotada").font = Font(bold=True); r += 1
last = sheet_table(ws, ["Ordem recebida", "Arquivo", "Documento", "Natureza", "Data", "Observações"],
                   [list(d) for d in DOCS], [14, 38, 40, 45, 12, 70], "TabDocs", start_row=r)
ws.freeze_panes = None
r = last + 2
ws.cell(row=r, column=1, value="Hierarquia adotada (o RFP não define precedência; ver CT-09 / PE-01)").font = Font(bold=True); r += 1
for h in HIERARQUIA:
    ws.cell(row=r, column=1, value=h); r += 1
r += 1
ws.cell(row=r, column=1, value="Problemas encontrados nos insumos").font = Font(bold=True); r += 1
for p in PROBLEMAS_INSUMOS:
    ws.cell(row=r, column=1, value="• " + p); r += 1
r += 1
ws.cell(row=r, column=1, value="Legenda de risco").font = Font(bold=True); r += 1
for k, v in [("Alto", "pode eliminar a proposta, gerar perda relevante ou exige decisão da diretoria"), ("Médio", "impacto em custo/prazo administrável com premissa ou ressalva"), ("Baixo", "requisito usual / formal")]:
    c = ws.cell(row=r, column=1, value=f"{k}: {v}"); c.fill = PatternFill("solid", fgColor=RISK_FILL[k]); r += 1
ws.column_dimensions["A"].width = 14

# ---------- Resumo ----------
ws = wb.create_sheet("Resumo")
ws["A1"] = "Principais riscos e pontos de atenção (top 10)"; ws["A1"].font = Font(bold=True, size=13, color="1F3864")
sheet_table(ws, ["#", "Ref. matriz", "Ponto"], [[i, a, b] for i, (a, b) in enumerate(TOP_RISCOS, 1)], [5, 14, 120], "TabTop", start_row=3)
from collections import Counter
start = 3 + len(TOP_RISCOS) + 3
ws.cell(row=start, column=1, value="Contagem da matriz por categoria e risco").font = Font(bold=True)
cats = list(dict.fromkeys(m["Categoria"] for m in M))
rows = []
for cat in cats:
    rows.append([cat] + [f'=COUNTIFS(Matriz!$B:$B,$A{start+2+len(rows)},Matriz!$J:$J,"{k}")' for k in ("Alto", "Médio", "Baixo")]
                + [f'=COUNTIF(Matriz!$B:$B,$A{start+2+len(rows)})'])
hdr_r = start + 1
for j, h in enumerate(["Categoria", "Alto", "Médio", "Baixo", "Total"], 1):
    c = ws.cell(row=hdr_r, column=j, value=h); c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
for i, rw in enumerate(rows, hdr_r + 1):
    for j, v in enumerate(rw, 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BORDER
tot = hdr_r + len(rows) + 1
ws.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
for j in range(2, 6):
    L = get_column_letter(j)
    c = ws.cell(row=tot, column=j, value=f"=SUM({L}{hdr_r+1}:{L}{tot-1})"); c.font = Font(bold=True); c.border = BORDER
ws.cell(row=tot + 2, column=1, value="Itens com impacto em Pricing:")
ws.cell(row=tot + 2, column=3, value='=COUNTIF(Matriz!$K:$K,"Sim")')
ws.cell(row=tot + 3, column=1, value="Itens com impacto Jurídico:")
ws.cell(row=tot + 3, column=3, value='=COUNTIF(Matriz!$L:$L,"Sim")')
ws.column_dimensions["A"].width = 22

# ---------- Matriz ----------
ws = wb.create_sheet("Matriz")
cols = COLS + ["Status interno", "Responsável", "Comentários das áreas"]
rows = [[m[c] for c in COLS] + ["Aberto", "", ""] for m in M]
widths = [8, 16, 30, 22, 26, 60, 20, 34, 28, 9, 10, 10, 22, 60, 60, 10, 14, 16, 30]
last = sheet_table(ws, cols, rows, widths, "TabMatriz", risk_col=10)
dv = DataValidation(type="list", formula1='"Aberto,Em análise,Resolvido,Aceito como risco"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"Q2:Q{last}")
for rr in range(2, last + 1):
    for col in ("K", "L"):
        c = ws[f"{col}{rr}"]
        if c.value == "Sim":
            c.font = Font(bold=True, color="C00000")

# ---------- Jurídico / Pricing ----------
def recorte(nome, chave, tab):
    w = wb.create_sheet(nome)
    sel = [m for m in M if m[chave] == "Sim"]
    cs = ["ID", "Categoria", "Tema", "Fonte principal", "Outras fontes", "Requisito / conteúdo (resumo)", "Posição vigente (após Q&A)",
          "Conflito / ambiguidade", "Risco", "Análise", "Ação recomendada", "Pergunta (ref.)"]
    order = {"Alto": 0, "Médio": 1, "Baixo": 2}
    sel.sort(key=lambda m: order[m["Risco"]])
    sheet_table(w, cs, [[m[c] for c in cs] for m in sel], [8, 16, 30, 20, 24, 55, 32, 26, 9, 60, 60, 10], tab, risk_col=9)
recorte("Jurídico", "Impacto Jurídico", "TabJuridico")
recorte("Pricing", "Impacto Pricing", "TabPricing")

# ---------- Conflitos ----------
ws = wb.create_sheet("Conflitos")
sheet_table(ws, ["ID", "Tema", "Fonte A", "Fonte B", "Fonte C / contexto", "Posição recomendada", "Gravidade", "Ref. matriz", "Pergunta"],
            [list(c) for c in CONFLITOS], [7, 26, 40, 40, 38, 50, 14, 16, 10], "TabConflitos")

# ---------- Esclarecimentos ----------
ws = wb.create_sheet("Esclarecimentos")
sheet_table(ws, ["ID", "Referência (documento/item)", "Pergunta proposta (texto para envio)", "Justificativa interna", "Prioridade", "Área", "Status", "Resposta do banco"],
            [list(p) + ["A enviar", ""] for p in PERGUNTAS], [7, 30, 90, 40, 10, 20, 12, 40], "TabPerguntas")

# ---------- Cronograma ----------
ws = wb.create_sheet("Cronograma")
sheet_table(ws, ["Data", "Evento", "Fonte", "Observação"], [list(c) for c in CRONO], [24, 70, 18, 45], "TabCrono")

# ---------- Pontuação ----------
ws = wb.create_sheet("Pontuação")
ws["A1"] = "Simulação da Nota Técnica (RFP 6.2/6.3 + regra oculta da aba 'Pesos')"; ws["A1"].font = Font(bold=True, size=12, color="1F3864")
hdr = ["Critério (RFP 6.2)", "Pontuação máxima", "Nota estimada (preencher)"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=3, column=j, value=h); c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER
crit = [("Metodologia de operação do SOC", 20, 16), ("Qualificação da equipe", 20, 16), ("Ferramentas e tecnologia", 15, 12), ("Plano de transição", 10, 8)]
for i, (a, b, cnota) in enumerate(crit, 4):
    for j, v in enumerate((a, b, cnota), 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BORDER
    ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor="FFF2CC")
ws["A8"] = "Soma dos critérios"; ws["B8"] = "=SUM(B4:B7)"; ws["C8"] = "=SUM(C4:C7)"
ws["A9"] = "Máximo declarado no RFP 6.2"; ws["B9"] = 60
ws["A10"] = "Diferença (inconsistência)"; ws["B10"] = "=B8-B9"
ws["A12"] = "Penalidade oculta: nº de não obrigatórios respondidos 'Não atende' (QT-02, QT-05)"; ws["B12"] = 0
ws["B12"].fill = PatternFill("solid", fgColor="FFF2CC")
ws["A13"] = "Pontos por item (aba 'Pesos')"; ws["B13"] = 2
ws["A14"] = "Nota técnica estimada após penalidade"; ws["B14"] = "=C8-B12*B13"
ws["A16"] = "Nota de corte (70%) com base 60"; ws["B16"] = "=0.7*B9"
ws["A17"] = "Nota de corte (70%) com base 65"; ws["B17"] = "=0.7*B8"
ws["A18"] = "Situação no cenário mais exigente (base 65)"; ws["B18"] = '=IF(B14>=B17,"Acima do corte","DESCLASSIFICA")'
ws["A19"] = "Folga sobre o corte mais exigente (pts)"; ws["B19"] = "=B14-B17"
ws["A21"] = "Composição da nota final (RFP 6.1): Técnica 60% + Preço 40%. A fórmula da Nota de Preço não foi informada (PE-07)."
ws["A22"] = "Células amarelas são entradas editáveis."
for rr in (8, 14, 18):
    ws[f"A{rr}"].font = Font(bold=True); ws[f"B{rr}"].font = Font(bold=True)
ws.column_dimensions["A"].width = 75; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 24

# ---------- Questionário ----------
ws = wb.create_sheet("Questionário")
sheet_table(ws, ["ID", "Requisito", "Obrigatório", "Peso (aba oculta)", "Penalidade oculta", "Resposta recomendada", "Justificativa", "Ref. matriz"],
            [list(q) for q in QUEST], [8, 55, 11, 12, 18, 50, 55, 14], "TabQuest")

for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
xlsx_path = os.path.join(OUT, "Matriz_Analise_RFP-2026-017_Banco_Horizonte.xlsx")
wb.save(xlsx_path)

# ======================= DOCX =======================
d = docx.Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
sec = d.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.8))

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color); tcPr.append(s)

def table(headers, rows, widths=None, risk_idx=None, size=8.5):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ""; p = c.paragraphs[0]; run = p.add_run(h); run.bold = True; run.font.size = Pt(size); run.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1F3864")
    for r in rows:
        cells = t.add_row().cells
        for j, v in enumerate(r):
            cells[j].text = ""; run = cells[j].paragraphs[0].add_run(str(v)); run.font.size = Pt(size)
            if risk_idx is not None and j == risk_idx and v in RISK_FILL:
                shade(cells[j], RISK_FILL[v])
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    d.add_paragraph()
    return t

d.add_heading("Matriz de Análise — RFP nº 2026-017", 0)
d.add_paragraph("Banco Horizonte Fictício S.A. — Serviços gerenciados de SOC 24x7x365 e IAM\n"
                "Data da análise: 18/10/2026 · Entrega das propostas: 27/11/2026 até 18h (prorrogada pelo Q&A #1) · Perguntas até 05/11/2026\n"
                "Documentos de origem fictícios para teste. Versão filtrável em Excel: Matriz_Analise_RFP-2026-017_Banco_Horizonte.xlsx")

d.add_heading("1. Resumo executivo", 1)
d.add_paragraph(
    "Objeto: SOC gerenciado 24x7x365, com SIEM fornecido pelo proponente e resposta a incidentes, mais gestão de identidades e acessos, "
    "por 36 meses, reajuste IPCA, nota final 60% técnica e 40% preço. Início previsto em 01/02/2027, colado ao fim do SIEM legado (31/01/2027, sem prorrogação).")
d.add_paragraph(
    "Leitura geral: a oportunidade é tecnicamente atendível, mas o RFP traz um nível de risco contratual e de indefinição de escopo acima do usual. "
    "Os itens 7.4 e 7.5 (multa sem teto e responsabilidade ilimitada) devem ser tratados como decisão de go/no-go pela diretoria, com o jurídico, antes do investimento na proposta. "
    "A volumetria só será informada ao vencedor, e há quatro números diferentes de usuários. Por isso o preço precisa ser montado com premissas explícitas e preços unitários de excedente.")
p = d.add_paragraph(); p.add_run("Recomendação: ").bold = True
p.add_run("seguir com a proposta de forma condicionada. (1) Enviar as 17 perguntas da seção 8 até 05/11/2026. (2) Obter o Anexo IV e confirmar se existe o Anexo II. "
          "(3) Levar CT-04 e CT-05 à diretoria. (4) Precificar com base de 5.000 usuários, faixa de EPS declarada, SIEM 100% no Brasil, SOAR e threat intel incluídos, "
          "e PAM, migração de logs e subsidiárias como opcionais.")
d.add_heading("Top 10 pontos de atenção", 2)
table(["#", "Ref.", "Ponto"], [[i, a, b] for i, (a, b) in enumerate(TOP_RISCOS, 1)], [1, 2.2, 21])

d.add_heading("2. Documentos analisados e hierarquia", 1)
d.add_paragraph("Os arquivos foram recebidos fora da ordem lógica (Q&A primeiro, RFP por último). A análise foi reordenada pela hierarquia abaixo, "
                "aplicando o Q&A sobre o RFP e o Anexo I. O RFP não tem cláusula de precedência (ver CT-09 / PE-01).")
table(["Ordem recebida", "Arquivo", "Documento", "Natureza", "Data", "Observações"], [list(x) for x in DOCS], [2, 5, 4.5, 5, 2, 7])
for h in HIERARQUIA:
    d.add_paragraph(h, style="List Bullet")
d.add_heading("Problemas encontrados nos insumos", 2)
for pr in PROBLEMAS_INSUMOS:
    d.add_paragraph(pr, style="List Bullet")

d.add_heading("3. Cronograma", 1)
table(["Data", "Evento", "Fonte", "Observação"], [list(c) for c in CRONO], [4, 10, 3.5, 7])

d.add_heading("4. Conflitos entre documentos", 1)
table(["ID", "Tema", "Fonte A", "Fonte B", "Fonte C / contexto", "Posição recomendada", "Gravidade", "Pergunta"],
      [[c[0], c[1], c[2], c[3], c[4], c[5], c[6], c[8]] for c in CONFLITOS], [1.2, 2.6, 4, 4, 3.8, 5.2, 1.8, 1.4])

d.add_heading("5. Matriz analítica completa", 1)
d.add_paragraph("Visão condensada. A análise detalhada e as colunas de impacto (Pricing/Jurídico) estão nas seções 6 e 7 e na planilha Excel.")
cats = list(dict.fromkeys(m["Categoria"] for m in M))
for cat in cats:
    d.add_heading(cat, 2)
    table(["ID", "Tema", "Fonte", "Requisito / conteúdo", "Tipo", "Posição vigente", "Risco", "Ação recomendada"],
          [[m["ID"], m["Tema"], m["Fonte principal"] + ("; " + m["Outras fontes"] if m["Outras fontes"] else ""), m["Requisito / conteúdo (resumo)"],
            m["Tipo"], m["Posição vigente (após Q&A)"], m["Risco"], m["Ação recomendada"]] for m in M if m["Categoria"] == cat],
          [1.3, 2.6, 2.8, 5.5, 2.2, 3.2, 1.3, 5.2], risk_idx=6, size=8)

order = {"Alto": 0, "Médio": 1, "Baixo": 2}
d.add_heading("6. Visão Jurídico", 1)
d.add_paragraph("Itens com impacto jurídico, do maior para o menor risco. Prioridades: CT-05 (responsabilidade ilimitada), CT-04 (multa sem teto), "
                "PR-06/TS-02 (declaração do Anexo III x nuvem em região livre), CT-07 (subcontratação x nuvem), CT-09 (precedência) e DM-02 (subsidiárias como partes).")
sel = sorted([m for m in M if m["Impacto Jurídico"] == "Sim"], key=lambda m: order[m["Risco"]])
table(["ID", "Tema", "Fonte", "Risco", "Análise", "Ação recomendada"],
      [[m["ID"], m["Tema"], m["Fonte principal"], m["Risco"], m["Análise"], m["Ação recomendada"]] for m in sel], [1.3, 3, 2.6, 1.3, 8.5, 7.5], risk_idx=3, size=8)

d.add_heading("7. Visão Pricing", 1)
d.add_paragraph("Premissas recomendadas para o modelo de preço:")
for t in [
    "Base de 5.000 usuários (Q&A #5), com preço unitário por usuário adicional e faixas até ~7.100 (6.200 ativos + ~900 das subsidiárias).",
    "SIEM fornecido pelo proponente, 100% em região de nuvem no Brasil, inclusive DR e backup (Anexo III / Res. CMN 4.893).",
    "Volumetria declarada como premissa (faixa de EPS e GB/dia) com preço por excedente. Base indicativa: 850 servidores, 2 DCs, 2 nuvens públicas, ~1.200 alertas/dia.",
    "Armazenamento em camadas: 12 meses online + 5 anos em arquivo, com preço por TB excedente.",
    "Incluídos no preço base: SOAR e threat intelligence (evitam até -4 pts de penalidade oculta), coordenador CISSP exclusivo, N1 remoto no Brasil (Q&A #6), transição e operação paralela sem custo adicional.",
    "Opcionais cotados à parte: PAM (desejável), migração de 24 meses de logs históricos, subsidiárias, visão executiva/GRC e forense.",
    "Fatores financeiros: pagamento em 60 dias após o aceite (custo financeiro), IPCA x componentes em USD (câmbio), prêmio de seguro cyber e provisão para multas de SLA sem teto.",
    "A estrutura final depende do Anexo IV (Planilha de Preços), que não foi recebido.",
]:
    d.add_paragraph(t, style="List Bullet")
sel = sorted([m for m in M if m["Impacto Pricing"] == "Sim"], key=lambda m: order[m["Risco"]])
table(["ID", "Tema", "Fonte", "Risco", "Posição vigente", "Ação recomendada"],
      [[m["ID"], m["Tema"], m["Fonte principal"], m["Risco"], m["Posição vigente (após Q&A)"], m["Ação recomendada"]] for m in sel], [1.3, 3.5, 2.6, 1.3, 5.5, 10], risk_idx=3, size=8)

d.add_heading("8. Perguntas de esclarecimento (enviar até 05/11/2026)", 1)
table(["ID", "Referência", "Pergunta proposta", "Prioridade", "Área"], [[p[0], p[1], p[2], p[4], p[5]] for p in PERGUNTAS], [1.3, 3.5, 15, 1.8, 2.8], size=8)
d.add_paragraph("Atenção (PE-08): a regra de penalidade foi encontrada em aba OCULTA da planilha. Antes de enviar, avaliar com o comercial e o jurídico a conveniência de citá-la diretamente.")

d.add_heading("9. Pontuação técnica e questionário", 1)
d.add_paragraph("RFP 6.2: Metodologia 20 + Equipe 20 + Ferramentas 15 + Transição 10 = 65 pontos, contra o máximo declarado de 60. "
                "Nota de corte de 70% (6.3): 42,0 pts (base 60) ou 45,5 pts (base 65). A aba oculta 'Pesos' acrescenta -2 pts por requisito não obrigatório "
                "respondido como 'Não atende' (QT-02 SOAR e QT-05 Threat Intel: até -4 pts). Recomendação: mirar em pelo menos 50 pts e responder 'Atende' em todos os QTs.")
table(["ID", "Requisito", "Obrig.", "Peso", "Penalidade", "Resposta recomendada", "Justificativa"],
      [list(q[:7]) for q in QUEST], [1.3, 5.5, 1.3, 1.2, 2.5, 5.5, 7])

d.add_heading("10. Checklist de habilitação e entrega", 1)
for t in [
    "Certificado ISO/IEC 27001 vigente, organismo acreditado, escopo incluindo o SOC (5.1a).",
    "Mínimo de 3 atestados de SOC para instituições financeiras com ≥2.000 usuários (5.1b).",
    "Comprovação de SOC em território nacional (5.1c; Q&A #2).",
    "Comprovação de equipe certificada, incluindo coordenador CISSP (5.2; Anexo I 5.1).",
    "Anexo III — Declaração de Localização de Dados assinada pelo representante legal (eliminatória).",
    "Questionário Técnico integralmente preenchido, colunas D e E de QT-01 a QT-06 (eliminatório).",
    "Proposta técnica em português, com no máximo 40 páginas excluídos os anexos.",
    "Proposta comercial no modelo do Anexo IV (a obter).",
    "Envio exclusivamente pelo portal de fornecedores até 27/11/2026, 18h.",
]:
    d.add_paragraph("☐ " + t)

d.add_heading("11. Limitações desta análise", 1)
for t in [
    "O Anexo III foi lido a partir de imagem escaneada (sem OCR). Conferir a transcrição com o original.",
    "O Anexo IV, o Anexo II e a minuta de contrato não foram recebidos.",
    "As referências à Res. CMN 4.893/2021 são gerais. O jurídico ou o compliance deve confirmar a redação vigente e as obrigações aplicáveis.",
    "Esta análise não substitui parecer jurídico.",
]:
    d.add_paragraph(t, style="List Bullet")

docx_path = os.path.join(OUT, "Matriz_Analise_RFP-2026-017_Banco_Horizonte.docx")
d.save(docx_path)
print(xlsx_path); print(docx_path); print("itens matriz:", len(M))
