#!/usr/bin/env python3
"""Exporta a matriz analítica (Markdown) para Excel.

Uso:
    python export_xlsx.py <matriz.md> [<saida.xlsx>]

Gera a aba "Matriz" (filtros, painel congelado, cor por criticidade) e a aba
"Resumo" (contagens por documento, tipo, criticidade e ação, e o status final).
"""
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WIDTHS = [10, 26, 16, 26, 22, 50, 24, 24, 60, 14, 12, 14, 14, 14, 40, 26, 30, 14, 22, 16, 22, 50, 22]
CRIT_FILL = {"Crítica": "F4B6B6", "Alta": "F8D7A8", "Média": "FFF2B3", "Baixa": "D5ECD4"}


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".xlsx")
    lines = src.read_text(encoding="utf-8").splitlines()
    header = next(l for l in lines if l.startswith("| ID |"))
    cols = [c.strip() for c in header.strip().strip("|").split("|")]
    rows = [[c.strip() for c in l.strip().strip("|").split("|")]
            for l in lines if re.match(r"^\|\s*D\d{2}-\d{3}\s*\|", l)]
    status = [l for l in lines if l.startswith(("✅", "⏸", "⛔", "Documentos processados", "Pendências", "Progresso"))]

    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz"
    ws.append(cols)
    head_fill = PatternFill("solid", fgColor="DDE3EA")
    for i, w in enumerate(WIDTHS[:len(cols)], 1):
        c = ws.cell(1, i)
        c.font, c.fill = Font(bold=True), head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ci = cols.index("Criticidade") if "Criticidade" in cols else None
    for r in rows:
        ws.append(r)
        row = ws.max_row
        for c in ws[row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if ci is not None:
            label = r[ci].split("·")[-1].strip() if "·" in r[ci] else ""
            if label in CRIT_FILL:
                ws.cell(row, ci + 1).fill = PatternFill("solid", fgColor=CRIT_FILL[label])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    rs = wb.create_sheet("Resumo")
    rs.column_dimensions["A"].width = 40
    rs.column_dimensions["B"].width = 12

    def block(title, counter):
        rs.append([title, "Itens"])
        rs.cell(rs.max_row, 1).font = Font(bold=True)
        rs.cell(rs.max_row, 2).font = Font(bold=True)
        for k, v in counter.most_common():
            rs.append([k, v])
        rs.append([])

    idx = {c: i for i, c in enumerate(cols)}
    block("Documento", Counter(r[idx["ID"]][:3] + " · " + r[idx["Documento"]] for r in rows))
    block("Tipo de item", Counter(r[idx["Tipo de item"]] for r in rows))
    block("Criticidade", Counter(r[idx["Criticidade"]].split("·")[-1].strip() for r in rows))
    block("Ação recomendada", Counter(r[idx["Ação recomendada"]] for r in rows))
    if status:
        rs.append(["Status"])
        rs.cell(rs.max_row, 1).font = Font(bold=True)
        for l in status:
            rs.append([l])
    wb.save(dst)
    print(f"{len(rows)} linhas exportadas para {dst}")


if __name__ == "__main__":
    main()
