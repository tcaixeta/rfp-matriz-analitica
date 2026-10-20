#!/usr/bin/env python3
"""Extrai o texto de documentos de concorrência com marcadores de localizador.

Uso:
    python extract_documents.py --out <pasta> arquivo1 arquivo2 ... [--start 1]

Os arquivos recebem os códigos D01, D02... na ordem em que são passados (use
--start para continuar a numeração numa análise incremental). Para cada
documento é gerado <pasta>/Dnn_<nome>.md com o texto marcado por localizador,
e <pasta>/inventario.md resume extensão e sinalizações (páginas com pouco texto,
abas ocultas, notas de slide, marcas de revisão...). Páginas de PDF com pouco
texto são salvas como PNG em <pasta>/imagens/ para leitura visual.

Formatos: .pdf, .docx, .pptx, .xlsx/.xlsm. Outros formatos são listados no
inventário como não processados.
"""
import argparse
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

LOW_TEXT_CHARS = 150          # corpo de página abaixo disso = possível imagem
PAGE_NUM_RE = re.compile(r"^(?:p[áa]gina|p[áa]g\.?|page)\s*(\d{1,4})(?:\s*(?:de|of|/)\s*\d+)?$|^(\d{1,4})(?:\s*(?:de|of|/)\s*\d+)?$",
                         re.IGNORECASE)


def _w(lines, text=""):
    lines.append(text)


# ----------------------------------------------------------------------------- PDF
def extract_pdf(path, code, out_dir, img_dir):
    import pdfplumber

    lines, flags = [], []
    with pdfplumber.open(path) as pdf:
        pages = [(p.extract_text() or "").splitlines() for p in pdf.pages]
        n = len(pages)
        # Linhas repetidas em metade ou mais das páginas = cabeçalho/rodapé
        counts = Counter(l.strip() for pg in pages for l in set(pg) if l.strip())
        repeated = {l for l, c in counts.items() if n >= 3 and c >= max(2, n // 2)}
        for i, (page, raw) in enumerate(zip(pdf.pages, pages), 1):
            printed = None
            for cand in raw[-3:][::-1] + raw[:2]:
                m = PAGE_NUM_RE.match(cand.strip())
                if m:
                    printed = m.group(1) or m.group(2)
                    break
            body = [l for l in raw if l.strip() and l.strip() not in repeated
                    and not PAGE_NUM_RE.match(l.strip())]
            body_chars = sum(len(l) for l in body)
            img_area = sum((im["x1"] - im["x0"]) * (im["bottom"] - im["top"]) for im in page.images)
            img_share = img_area / (page.width * page.height) if page.width and page.height else 0
            loc = f"p. {i}" + (f" (impressa: {printed})" if printed else " (impressa: sem número)")
            _w(lines, f"## [{loc}]")
            if body_chars < LOW_TEXT_CHARS:
                png = img_dir / f"{code}_p{i}.png"
                page.to_image(resolution=150).save(str(png))
                note = (f"POSSÍVEL CONTEÚDO EM IMAGEM NÃO LIDO — corpo com {body_chars} caracteres"
                        f"{f', imagem ocupa {img_share:.0%} da página' if img_share else ''}. "
                        f"Abra {png.name} e leia o conteúdo; marque o localizador com (OCR).")
                flags.append(f"{loc}: {note}")
                _w(lines, f"> ⚠ {note}")
            elif img_share > 0.3:
                flags.append(f"{loc}: imagem ocupa {img_share:.0%} da página — confira se há texto na imagem.")
                _w(lines, f"> ⚠ Imagem ocupa {img_share:.0%} da página; confira se há texto nela.")
            for l in raw:
                _w(lines, l)
            _w(lines)
        if repeated:
            flags.append("Cabeçalho/rodapé repetido (ignorado na contagem de texto): " + " · ".join(sorted(repeated))[:200])
    return lines, f"{n} páginas", flags


# ----------------------------------------------------------------------------- DOCX
def extract_docx(path, code, out_dir, img_dir):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    lines, flags = [], []
    heading, n_par, n_tab = "início do documento", 0, 0
    _w(lines, f"## [§ {heading}]")
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if re.match(r"(heading|título|titulo)\s*\d", style, re.IGNORECASE) or style.lower() in ("title", "título"):
                heading = text
                _w(lines)
                _w(lines, f"## [§ {heading}]")
                continue
            n_par += 1
            _w(lines, text)
        elif tag == "tbl":
            n_tab += 1
            t = Table(el, doc)
            _w(lines)
            _w(lines, f"### [§ {heading} · tabela {n_tab}]")
            for r, row in enumerate(t.rows, 1):
                cells = [c.text.strip().replace("|", "/") for c in row.cells]
                _w(lines, f"| L{r} | " + " | ".join(cells) + " |")
            _w(lines)
    # Partes que o texto corrido não mostra
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
        if "<w:ins " in xml or "<w:del " in xml:
            flags.append("Há marcas de revisão (inserções/exclusões não aceitas): confira a versão vigente do texto.")
        if "<w:vanish/>" in xml:
            flags.append("Há texto oculto (formatação 'oculto') no documento.")
        if "word/comments.xml" in names:
            cx = z.read("word/comments.xml").decode("utf-8", "ignore")
            comments = [re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<w:comment\b.*?</w:comment>", cx, re.S)]
            comments = [c for c in comments if c]
            if comments:
                flags.append(f"{len(comments)} comentário(s) no documento — veja a seção [comentários].")
                _w(lines, "## [comentários do documento]")
                for i, c in enumerate(comments, 1):
                    _w(lines, f"- comentário {i}: {c}")
        hf = [n for n in names if re.match(r"word/(header|footer)\d*\.xml", n)]
        hf_text = []
        for n in hf:
            t = re.sub(r"<[^>]+>", " ", z.read(n).decode("utf-8", "ignore"))
            t = re.sub(r"\s+", " ", t).strip()
            if t:
                hf_text.append(t)
        if hf_text:
            _w(lines, "## [cabeçalhos e rodapés]")
            for t in sorted(set(hf_text)):
                _w(lines, f"- {t}")
    flags.append("Word não tem paginação fixa: use seção/cláusula como localizador; página, se informada, como 'aprox.'.")
    return lines, f"{n_par} parágrafos, {n_tab} tabela(s)", flags


# ----------------------------------------------------------------------------- PPTX
def _shape_texts(shapes):
    for sh in shapes:
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # grupo
            yield from _shape_texts(sh.shapes)
        elif getattr(sh, "has_table", False) and sh.has_table:
            for r, row in enumerate(sh.table.rows, 1):
                yield "| L%d | " % r + " | ".join(c.text.strip().replace("|", "/") for c in row.cells) + " |"
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    yield t


def extract_pptx(path, code, out_dir, img_dir):
    from pptx import Presentation

    prs = Presentation(path)
    lines, flags = [], []
    for i, slide in enumerate(prs.slides, 1):
        hidden = slide._element.get("show") == "0"
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip()
        texts = list(_shape_texts(slide.shapes))
        if not title and texts:
            title = texts[0]
        pics = sum(1 for sh in slide.shapes if sh.shape_type == 13)
        label = f'slide {i} — "{title}"' + (" (OCULTO)" if hidden else "")
        _w(lines, f"## [{label}]")
        if hidden:
            flags.append(f"slide {i} está oculto na apresentação — analise o conteúdo mesmo assim.")
        if pics:
            _w(lines, f"> {pics} imagem(ns) no slide; confira se contêm texto relevante.")
        for t in texts:
            if t != title:
                _w(lines, t)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                flags.append(f"slide {i} tem notas do apresentador — podem conter escopo não exibido.")
                _w(lines, f"### [slide {i} — notas do apresentador]")
                _w(lines, notes)
        _w(lines)
    return lines, f"{len(prs.slides)} slides", flags


# ----------------------------------------------------------------------------- XLSX
def extract_xlsx(path, code, out_dir, img_dir):
    from openpyxl import load_workbook

    wf = load_workbook(path, data_only=False)
    wv = load_workbook(path, data_only=True)
    lines, flags, sheets = [], [], []
    for ws in wf.worksheets:
        vs = wv[ws.title]
        state = ws.sheet_state
        sheets.append(f'"{ws.title}"' + (f" ({'OCULTA' if state == 'hidden' else 'MUITO OCULTA'})" if state != "visible" else ""))
        if state != "visible":
            flags.append(f'aba "{ws.title}" está {"oculta" if state == "hidden" else "muito oculta (veryHidden)"} — analise o conteúdo.')
        hid_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
        hid_cols = [c for c, d in ws.column_dimensions.items() if d.hidden]
        if hid_rows or hid_cols:
            flags.append(f'aba "{ws.title}": linhas ocultas {hid_rows or "—"} · colunas ocultas {hid_cols or "—"}.')
        _w(lines, f'## [aba "{ws.title}"' + (" — OCULTA" if state != "visible" else "") + f" · {ws.dimensions}]")
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    cached = vs[c.coordinate].value
                    v = f"{v}  → valor: {cached}"
                text = str(v).replace("\n", " ")
                mark = " (linha oculta)" if c.row in hid_rows else ""
                _w(lines, f'[aba "{ws.title}" · {c.coordinate}]{mark} {text}')
                if c.comment:
                    _w(lines, f'[aba "{ws.title}" · {c.coordinate} · comentário] {c.comment.text.strip()}')
        _w(lines)
    return lines, f"{len(wf.worksheets)} abas: " + ", ".join(sheets), flags


HANDLERS = {".pdf": extract_pdf, ".docx": extract_docx, ".pptx": extract_pptx,
            ".xlsx": extract_xlsx, ".xlsm": extract_xlsx}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--start", type=int, default=1, help="número do primeiro código (padrão: 1 → D01)")
    a = ap.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    img_dir = a.out / "imagens"
    img_dir.mkdir(exist_ok=True)

    inv = ["# Inventário da extração", "",
           "| Código | Arquivo | Formato | Extensão | Arquivo extraído |", "|---|---|---|---|---|"]
    details = []
    for k, f in enumerate(a.files, a.start):
        code = f"D{k:02d}"
        ext = f.suffix.lower()
        if not f.exists():
            inv.append(f"| {code} | {f.name} | — | ARQUIVO NÃO ENCONTRADO | — |")
            details.append((code, f.name, ["Arquivo não encontrado: registre como Registro de cobertura (arquivo listado e não anexado)."]))
            continue
        if ext not in HANDLERS:
            inv.append(f"| {code} | {f.name} | {ext} | formato não suportado pelo script | — |")
            details.append((code, f.name, [f"Formato {ext} não suportado: leia com outra ferramenta ou registre como Registro de cobertura."]))
            continue
        try:
            lines, extent, flags = HANDLERS[ext](f, code, a.out, img_dir)
        except Exception as e:  # arquivo corrompido, protegido etc.
            inv.append(f"| {code} | {f.name} | {ext} | ERRO NA LEITURA | — |")
            details.append((code, f.name, [f"Erro ao ler ({type(e).__name__}: {e}). Registre como Registro de cobertura."]))
            continue
        out_name = f"{code}_{re.sub(r'[^A-Za-z0-9_.-]+', '_', f.stem)}.md"
        (a.out / out_name).write_text(f"# {code} — {f.name}\n\n" + "\n".join(lines) + "\n", encoding="utf-8")
        inv.append(f"| {code} | {f.name} | {ext} | {extent} | {out_name} |")
        details.append((code, f.name, flags))

    inv += ["", "## Sinalizações", ""]
    for code, name, flags in details:
        inv.append(f"### {code} — {name}")
        inv += [f"- {fl}" for fl in flags] or ["- nenhuma"]
        inv.append("")
    (a.out / "inventario.md").write_text("\n".join(inv) + "\n", encoding="utf-8")
    if not any(img_dir.iterdir()):
        img_dir.rmdir()
    print((a.out / "inventario.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    sys.exit(main())
